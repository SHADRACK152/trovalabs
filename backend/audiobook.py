"""
Audiobook Service - Convert novels/stories to audiobooks
Supports PDF, TXT, EPUB, DOCX files
"""
import logging
import os
import hashlib
import json
from pathlib import Path
from typing import Optional, List, Dict
from title_generator import get_title_generator
from datetime import datetime
import asyncio
import io

logger = logging.getLogger(__name__)


class AudiobookService:
    """Service for converting documents to audiobooks"""
    
    def __init__(self):
        """Initialize audiobook service"""
        self.audiobooks_dir = Path("audiobooks")
        self.audiobooks_dir.mkdir(exist_ok=True)
        
        # Create subdirectories
        self.uploads_dir = self.audiobooks_dir / "uploads"
        self.audio_dir = self.audiobooks_dir / "audio"
        self.metadata_dir = self.audiobooks_dir / "metadata"
        
        for directory in [self.uploads_dir, self.audio_dir, self.metadata_dir]:
            directory.mkdir(exist_ok=True)
        
        logger.info("📚 Audiobook Service initialized")
    
    async def extract_text_from_file(self, file_bytes: bytes, filename: str) -> str:
        """
        Extract text from various file formats
        
        Args:
            file_bytes: File content
            filename: Original filename
            
        Returns:
            Extracted text content
        """
        file_ext = Path(filename).suffix.lower()
        
        try:
            if file_ext == '.txt':
                return await self._extract_from_txt(file_bytes)
            elif file_ext == '.pdf':
                return await self._extract_from_pdf(file_bytes)
            elif file_ext == '.epub':
                return await self._extract_from_epub(file_bytes)
            elif file_ext in ['.docx', '.doc']:
                return await self._extract_from_docx(file_bytes)
            else:
                raise ValueError(f"Unsupported file format: {file_ext}")
        
        except Exception as e:
            logger.error(f"Text extraction error: {e}")
            raise
    
    async def _extract_from_txt(self, file_bytes: bytes) -> str:
        """Extract text from TXT file"""
        try:
            # Try UTF-8 first, then fallback encodings
            for encoding in ['utf-8', 'latin-1', 'cp1252']:
                try:
                    return file_bytes.decode(encoding)
                except UnicodeDecodeError:
                    continue
            raise ValueError("Could not decode text file")
        except Exception as e:
            logger.error(f"TXT extraction error: {e}")
            raise
    
    async def _extract_from_pdf(self, file_bytes: bytes) -> str:
        """Extract text from PDF file - ALL pages"""
        try:
            import PyPDF2
            pdf_file = io.BytesIO(file_bytes)
            reader = PyPDF2.PdfReader(pdf_file)
            
            total_pages = len(reader.pages)
            logger.info(f"📄 Extracting text from PDF: {total_pages} pages")
            
            text = ""
            for page_num, page in enumerate(reader.pages, 1):
                page_text = page.extract_text()
                text += page_text + "\n"
                if page_num % 10 == 0:  # Log every 10 pages
                    logger.info(f"   Processed {page_num}/{total_pages} pages...")
            
            logger.info(f"✅ Extracted {len(text)} characters from {total_pages} pages")
            return text.strip()
        
        except ImportError:
            logger.error("PyPDF2 not installed. Install: pip install PyPDF2")
            raise ValueError("PDF support not available. Install PyPDF2.")
        except Exception as e:
            logger.error(f"PDF extraction error: {e}")
            raise
    
    async def _extract_from_epub(self, file_bytes: bytes) -> str:
        """Extract text from EPUB file - ALL chapters"""
        try:
            import ebooklib
            from ebooklib import epub
            from bs4 import BeautifulSoup
            
            book = epub.read_epub(io.BytesIO(file_bytes))
            
            documents = [item for item in book.get_items() 
                        if item.get_type() == ebooklib.ITEM_DOCUMENT]
            logger.info(f"📚 Extracting text from EPUB: {len(documents)} documents")
            
            text = ""
            for idx, item in enumerate(documents, 1):
                soup = BeautifulSoup(item.get_content(), 'html.parser')
                text += soup.get_text() + "\n"
                if idx % 5 == 0:  # Log every 5 documents
                    logger.info(f"   Processed {idx}/{len(documents)} documents...")
            
            logger.info(f"✅ Extracted {len(text)} characters from EPUB")
            return text.strip()
        
        except ImportError:
            logger.error("ebooklib not installed. Install: pip install ebooklib beautifulsoup4")
            raise ValueError("EPUB support not available. Install ebooklib.")
        except Exception as e:
            logger.error(f"EPUB extraction error: {e}")
            raise
    
    async def _extract_from_docx(self, file_bytes: bytes) -> str:
        """Extract text from DOCX file - ALL paragraphs"""
        try:
            import docx
            
            doc_file = io.BytesIO(file_bytes)
            doc = docx.Document(doc_file)
            
            total_paragraphs = len(doc.paragraphs)
            logger.info(f"📝 Extracting text from DOCX: {total_paragraphs} paragraphs")
            
            text = ""
            for para in doc.paragraphs:
                text += para.text + "\n"
            
            logger.info(f"✅ Extracted {len(text)} characters from DOCX")
            return text.strip()
        
        except ImportError:
            logger.error("python-docx not installed. Install: pip install python-docx")
            raise ValueError("DOCX support not available. Install python-docx.")
        except Exception as e:
            logger.error(f"DOCX extraction error: {e}")
            raise
    
    def split_into_chapters(self, text: str, max_chars: int = 5000) -> List[Dict[str, any]]:
        """
        Split text into chapters/segments for TTS processing
        Intelligently detects chapter markers and boundaries
        
        Args:
            text: Full text content
            max_chars: Maximum characters per chapter
            
        Returns:
            List of chapter dictionaries with text and metadata
        """
        chapters = []
        
        # Try to detect chapter markers with improved patterns
        lines = text.split('\n')
        current_chapter = []
        current_text = ""
        chapter_num = 1
        last_chapter_title = None
        
        # Common chapter patterns
        import re
        chapter_patterns = [
            r'^chapter\s+\d+',
            r'^chapter\s+[ivxlcdm]+',  # Roman numerals
            r'^part\s+\d+',
            r'^part\s+[ivxlcdm]+',
            r'^section\s+\d+',
            r'^\d+\.\s+[A-Z]',  # "1. Title"
            r'^[IVXLCDM]+\.\s+[A-Z]',  # "I. Title"
            r'^ch(?:apter)?\b',
            r'^\d+\)\s+[A-Z]',  # "1) Title"
        ]
        
        for line in lines:
            line_stripped = line.strip()
            
            if not line_stripped:
                current_text += "\n"
                continue
            
            # Check if this line is a chapter marker
            is_chapter_marker = False
            for pattern in chapter_patterns:
                if re.match(pattern, line_stripped, re.IGNORECASE):
                    is_chapter_marker = True
                    break
            
            # Also detect all-caps short lines as potential chapter titles
            if (not is_chapter_marker and 
                line_stripped.isupper() and 
                len(line_stripped) < 100 and 
                len(line_stripped) > 3 and
                len(current_text) > 500):
                is_chapter_marker = True
            
            # Special handling: if the line is a 'Chapter X: Title' pattern, extract the title after colon
            if not is_chapter_marker:
                m = re.match(r'^chapter\s+(\d+)\s*[:.\-\)]?\s*(.*)$', line_stripped, re.IGNORECASE)
                if m:
                    is_chapter_marker = True
                    possible_title = m.group(2).strip()
                    if possible_title:
                        last_chapter_title = possible_title

            if is_chapter_marker and current_text:
                # Save previous chapter
                chapters.append({
                    "chapter_number": chapter_num,
                    "title": last_chapter_title or f"Chapter {chapter_num}",
                    "content": current_text.strip(),
                    "char_count": len(current_text)
                })
                last_chapter_title = line_stripped
                current_chapter = [line_stripped]
                current_text = ""
                chapter_num += 1
            else:
                current_chapter.append(line)
                current_text += line + "\n"
                
                # Split if too long (but try to split at paragraph boundaries)
                if len(current_text) >= max_chars:
                    # Try to find last paragraph break
                    split_pos = current_text.rfind('\n\n', 0, max_chars)
                    if split_pos > max_chars * 0.7:  # At least 70% through
                        chapter_text = current_text[:split_pos].strip()
                        current_text = current_text[split_pos:].strip()
                    else:
                        chapter_text = current_text.strip()
                        current_text = ""
                    
                    chapters.append({
                        "chapter_number": chapter_num,
                        "title": last_chapter_title or f"Chapter {chapter_num}",
                        "content": chapter_text,
                        "char_count": len(chapter_text)
                    })
                    last_chapter_title = None
                    current_chapter = []
                    chapter_num += 1
        
        # Add remaining text
        if current_text.strip():
            chapters.append({
                "chapter_number": chapter_num,
                "title": last_chapter_title or f"Chapter {chapter_num}",
                "content": current_text.strip(),
                "char_count": len(current_text)
            })
        
        # Merge very small chapters into previous chapters to avoid tiny fragments
        merged_chapters = []
        for ch in chapters:
            if merged_chapters and ch['char_count'] < 120:
                # Append small chapter to previous
                prev = merged_chapters[-1]
                prev['content'] = prev['content'] + "\n\n" + ch['content']
                prev['char_count'] = len(prev['content'])
            else:
                merged_chapters.append(ch)

        chapters = merged_chapters
        logger.info(f"📚 Split into {len(chapters)} chapters")
        for ch in chapters[:5]:  # Log first 5 chapter titles
            logger.info(f"  - {ch['title'][:80]}")
        
        return chapters
    
    async def create_audiobook(
        self,
        file_bytes: bytes,
        filename: str,
        user_id: str,
        title: Optional[str] = None,
        author: Optional[str] = None,
        voice: str = "lisa",
        tts_provider: str = "gtts"
    ) -> Dict[str, any]:
        """
        Create audiobook from uploaded file
        
        Args:
            file_bytes: Uploaded file content
            filename: Original filename
            user_id: User ID
            title: Book title (optional)
            author: Author name (optional)
            voice: TTS voice to use
            tts_provider: TTS provider (gtts, elevenlabs, etc.)
            
        Returns:
            Audiobook metadata
        """
        try:
            logger.info(f"Creating audiobook from {filename} for user {user_id}")
            
            # Generate audiobook ID
            audiobook_id = hashlib.md5(
                f"{user_id}_{filename}_{datetime.now().isoformat()}".encode()
            ).hexdigest()[:16]
            
            # Save uploaded file
            upload_path = self.uploads_dir / f"{audiobook_id}_{filename}"
            with open(upload_path, 'wb') as f:
                f.write(file_bytes)
            
            # Extract text
            logger.info("Extracting text...")
            text = await self.extract_text_from_file(file_bytes, filename)
            word_count = len(text.split())
            
            # Split into chapters
            logger.info("Splitting into chapters...")
            chapters = self.split_into_chapters(text)

            # Enrich chapter titles using LLM where appropriate
            try:
                title_gen = get_title_generator()
                # Use a short prompt per chapter to generate a short heading
                for idx, ch in enumerate(chapters):
                    current_title = ch.get('title') or ''
                    if current_title and not current_title.lower().startswith('chapter') and not current_title.lower().startswith('part'):
                        continue  # already a meaningful title
                    excerpt = ch.get('content', '')[:600]
                    try:
                        new_title = await title_gen.generate_title(excerpt, ai_response=None, conversation_context=None)
                        # sanitize and update
                        if new_title and len(new_title.strip()) > 0:
                            chapters[idx]['title'] = new_title
                    except Exception:
                        # ignore and keep existing title
                        continue
            except Exception:
                # If title generator not available, skip
                pass
            
            # Create metadata
            metadata = {
                "audiobook_id": audiobook_id,
                "user_id": user_id,
                "title": title or Path(filename).stem,
                "author": author or "Unknown",
                "original_filename": filename,
                "upload_path": str(upload_path),
                "voice": voice,
                "tts_provider": tts_provider,
                "word_count": word_count,
                "character_count": len(text),
                "chapter_count": len(chapters),
                "chapters": chapters,
                "created_at": datetime.now().isoformat(),
                "status": "ready",  # ready, processing, completed, failed
                "audio_files": []
            }
            
            # Save metadata
            metadata_path = self.metadata_dir / f"{audiobook_id}.json"
            with open(metadata_path, 'w', encoding='utf-8') as f:
                json.dump(metadata, f, indent=2)
            
            logger.info(f"✅ Audiobook created: {audiobook_id}")
            return metadata
        
        except Exception as e:
            logger.error(f"❌ Audiobook creation error: {e}")
            raise
    
    async def generate_audio_for_chapter(
        self,
        audiobook_id: str,
        chapter_number: int,
        tts_service,
        voice_override: Optional[str] = None,
    ) -> str:
        """
        Generate audio for a specific chapter with introduction
        
        Args:
            audiobook_id: Audiobook ID
            chapter_number: Chapter number
            tts_service: TTS service instance
            
        Returns:
            Path to generated audio file
        """
        try:
            # Load metadata
            metadata_path = self.metadata_dir / f"{audiobook_id}.json"
            with open(metadata_path, 'r', encoding='utf-8') as f:
                metadata = json.load(f)
            
            # Validate chapters exist
            if "chapters" not in metadata or not metadata["chapters"]:
                logger.error(f"❌ No chapters found in audiobook {audiobook_id}")
                raise ValueError("No chapters found in audiobook")
            
            # Validate chapter number
            if chapter_number < 1 or chapter_number > len(metadata["chapters"]):
                logger.error(f"❌ Invalid chapter number {chapter_number}, audiobook has {len(metadata['chapters'])} chapters")
                raise ValueError(f"Invalid chapter number {chapter_number}")
            
            # Get chapter
            chapter = metadata["chapters"][chapter_number - 1]
            chapter_title = chapter.get("title", f"Chapter {chapter_number}")
            text = chapter.get("content", chapter.get("text", ""))
            
            if not text:
                logger.error(f"❌ Chapter {chapter_number} has no content")
                raise ValueError(f"Chapter {chapter_number} has no content")
            
            logger.info(f"🎙️ Generating audio for: {chapter_title}")
            logger.info(f"   Text length: {len(text)} characters")
            
            # Create introduction text
            intro_text = f"{chapter_title}. "
            
            # Add book title on first chapter
            if chapter_number == 1:
                book_title = metadata.get("title", "")
                book_author = metadata.get("author", "")
                if book_title:
                    intro_text = f"{book_title}"
                    if book_author:
                        intro_text += f", by {book_author}. "
                    intro_text += f" {chapter_title}. "
            
            # Combine introduction with content
            full_text = intro_text + text
            
            # Limit text length for TTS
            if len(full_text) > 5000:
                logger.warning(f"Chapter too long, truncating to 5000 chars")
                # Keep introduction, truncate content
                available_chars = 5000 - len(intro_text) - 3
                text = text[:available_chars] + "..."
                full_text = intro_text + text
            
            try:
                # Determine voice selection from metadata or voice_override
                voice = voice_override if voice_override else (metadata.get("voice", "lisa") if metadata else "lisa")

                # If voice corresponds to free natural voices, delegate to that service
                try:
                    from elevenlabs_tts import get_free_tts_service
                    free_tts_service = get_free_tts_service()
                    free_voices = free_tts_service.VOICES.keys()
                except Exception:
                    free_voices = []

                # If using a browser-only voice (e.g., 'browser:Samantha'), fallback to 'lisa' for server-side generation
                if voice and voice.startswith('browser:'):
                    voice = 'lisa'

                if voice and voice.lower() in free_voices:
                    audio_bytes = await free_tts_service.synthesize(full_text, voice.lower())
                elif voice and voice.startswith("clone:"):
                    # Synthesise using user cloned voice
                    profile_name = (
                        voice.split(":", 1)[1] if ":" in voice else "default"
                    )
                    try:
                        from voice_clone import get_voice_clone_service
                        voice_clone = get_voice_clone_service()
                        audio_bytes = await voice_clone.synthesize_with_cloned_voice(
                            full_text,
                            user_id=metadata.get("user_id"),
                            profile_name=profile_name,
                        )
                    except Exception as vc_e:
                        logger.warning(
                            f"Cloned voice synthesis failed: {vc_e}; falling back to generic TTS"
                        )
                        audio_bytes = await tts_service.synthesize(
                            text=full_text,
                            tone="neutral",
                            voice="lisa"
                        )
                else:
                    # Use generic TTS service (gTTS / OpenAI) and pass voice if supported
                    audio_bytes = await tts_service.synthesize(
                        text=full_text,
                        tone="neutral",
                        voice=voice
                    )
                
                logger.info(f"✅ Generated {len(audio_bytes)} bytes")
            except Exception as tts_error:
                logger.error(f"TTS synthesis failed: {tts_error}")
                raise
            
                # Save audio file. If an override voice was used, include it in the filename to avoid collisions
            safe_voice = None
            if voice and voice != metadata.get("voice", "lisa"):
                # sanitize voice string for filename
                import re
                safe_voice = re.sub(r"[^a-zA-Z0-9_-]", "_", voice)

            if safe_voice:
                audio_filename = f"{audiobook_id}_chapter_{chapter_number:03d}_{safe_voice}.mp3"
            else:
                audio_filename = f"{audiobook_id}_chapter_{chapter_number:03d}.mp3"
            audio_path = self.audio_dir / audio_filename
            
            with open(audio_path, 'wb') as f:
                f.write(audio_bytes)
            
            logger.info(f"✅ Audio generated: {audio_filename} ({len(audio_bytes)} bytes)")
            return str(audio_path)
        
        except Exception as e:
            logger.error(f"❌ Audio generation error: {e}")
            raise
    
    async def process_audiobook(
        self,
        audiobook_id: str,
        tts_service,
        progress_callback=None
    ) -> Dict[str, any]:
        """
        Process entire audiobook (generate all chapter audio files)
        
        Args:
            audiobook_id: Audiobook ID
            tts_service: TTS service instance
            progress_callback: Optional callback for progress updates
            
        Returns:
            Updated metadata
        """
        try:
            # Load metadata
            metadata_path = self.metadata_dir / f"{audiobook_id}.json"
            with open(metadata_path, 'r', encoding='utf-8') as f:
                metadata = json.load(f)
            
            metadata["status"] = "processing"
            metadata["audio_files"] = []
            metadata["current_chapter"] = 0
            metadata["total_chapters_to_process"] = 0
            
            # Save initial status
            with open(metadata_path, 'w', encoding='utf-8') as f:
                json.dump(metadata, f, indent=2)
            
            # Total chapters count
            # (Use metadata['chapters'] directly when needed.)
            
            # Process ALL chapters (entire document)
            chapters_to_process = metadata["chapters"]
            
            metadata["total_chapters_to_process"] = len(chapters_to_process)
            
            logger.info(
                f"📚 Processing ALL {len(chapters_to_process)} chapters..."
            )
            
            # Generate audio for each chapter
            for idx, chapter in enumerate(chapters_to_process, 1):
                try:
                    logger.info(
                        f"Starting chapter {idx}/{len(chapters_to_process)}"
                    )
                    
                    # Update processing message
                    metadata["processing_message"] = (
                        f"Generating audio for chapter {idx}/{len(chapters_to_process)}: "
                        f"{chapter['title'][:50]}"
                    )
                    
                    # Save progress
                    with open(metadata_path, 'w', encoding='utf-8') as f:
                        json.dump(metadata, f, indent=2)
                    
                    audio_path = await self.generate_audio_for_chapter(
                        audiobook_id=audiobook_id,
                        chapter_number=idx,
                        tts_service=tts_service
                    )
                    
                    metadata["audio_files"].append({
                        "chapter_number": idx,
                        "title": chapter["title"],
                        "audio_path": audio_path,
                        "duration": None  # TODO: Calculate audio duration
                    })
                    
                    # Update current_chapter AFTER successful completion
                    metadata["current_chapter"] = idx
                    
                    # Save progress after each chapter
                    with open(metadata_path, 'w', encoding='utf-8') as f:
                        json.dump(metadata, f, indent=2)
                    
                    # Progress callback
                    if progress_callback:
                        await progress_callback(idx, len(chapters_to_process))
                    
                    logger.info(
                        f"✅ Progress: {idx}/{len(chapters_to_process)} "
                        f"chapters completed"
                    )
                
                except Exception as e:
                    logger.error(f"❌ Error processing chapter {idx}: {e}")
                    metadata["processing_message"] = (
                        f"Error in chapter {idx}: {str(e)}"
                    )
                    with open(metadata_path, 'w', encoding='utf-8') as f:
                        json.dump(metadata, f, indent=2)
                    # Continue with next chapter instead of failing completely
                    continue
            
            metadata["status"] = "completed"
            metadata["completed_at"] = datetime.now().isoformat()
            metadata["current_chapter"] = len(chapters_to_process)
            metadata["processing_message"] = "Processing complete!"
            
            # Save updated metadata
            with open(metadata_path, 'w', encoding='utf-8') as f:
                json.dump(metadata, f, indent=2)
            
            logger.info(f"✅ Audiobook processing completed: {audiobook_id}")
            return metadata
        
        except Exception as e:
            logger.error(f"❌ Audiobook processing error: {e}")
            
            # Update status to failed
            metadata["status"] = "failed"
            metadata["error"] = str(e)
            with open(metadata_path, 'w', encoding='utf-8') as f:
                json.dump(metadata, f, indent=2)
            
            raise
    
    def get_audiobook(self, audiobook_id: str) -> Optional[Dict[str, any]]:
        """Get audiobook metadata"""
        try:
            metadata_path = self.metadata_dir / f"{audiobook_id}.json"
            if not metadata_path.exists():
                return None
            
            with open(metadata_path, 'r', encoding='utf-8') as f:
                return json.load(f)
        
        except Exception as e:
            logger.error(f"Error loading audiobook: {e}")
            return None
    
    def get_user_audiobooks(self, user_id: str) -> List[Dict[str, any]]:
        """Get all audiobooks for a user"""
        audiobooks = []
        
        for metadata_file in self.metadata_dir.glob("*.json"):
            try:
                with open(metadata_file, 'r', encoding='utf-8') as f:
                    metadata = json.load(f)
                    
                if metadata.get("user_id") == user_id:
                    # Return summary (without full chapter text)
                    summary = {
                        k: v for k, v in metadata.items()
                        if k != "chapters"
                    }
                    summary["chapter_count"] = len(
                        metadata.get("chapters", [])
                    )
                    audiobooks.append(summary)
            
            except Exception as e:
                logger.error(
                    f"Error loading metadata file {metadata_file}: {e}"
                )
                continue
        
        # Sort by creation date (newest first)
        audiobooks.sort(key=lambda x: x.get("created_at", ""), reverse=True)
        
        return audiobooks
    
    def get_chapter_audio(
        self,
        audiobook_id: str,
        chapter_number: int,
        voice_override: Optional[str] = None,
    ) -> Optional[bytes]:
        """Get audio file for a specific chapter"""
        try:
            # If voice_override is set, try the voice-specific filename first
            if voice_override:
                import re
                safe_voice = re.sub(r"[^a-zA-Z0-9_-]", "_", voice_override)
                audio_filename_vo = f"{audiobook_id}_chapter_{chapter_number:03d}_{safe_voice}.mp3"
                audio_path_vo = self.audio_dir / audio_filename_vo
                if audio_path_vo.exists():
                    with open(audio_path_vo, 'rb') as f:
                        return f.read()

            audio_filename = f"{audiobook_id}_chapter_{chapter_number:03d}.mp3"
            audio_path = self.audio_dir / audio_filename
            
            if not audio_path.exists():
                return None
            
            with open(audio_path, 'rb') as f:
                return f.read()
        
        except Exception as e:
            logger.error(f"Error loading chapter audio: {e}")
            return None
    
    def delete_audiobook(self, audiobook_id: str, user_id: str) -> bool:
        """Delete an audiobook and all its files"""
        try:
            # Load metadata to verify ownership
            metadata = self.get_audiobook(audiobook_id)
            if not metadata or metadata.get("user_id") != user_id:
                return False
            
            # Delete audio files
            for audio_file in metadata.get("audio_files", []):
                audio_path = Path(audio_file["audio_path"])
                if audio_path.exists():
                    audio_path.unlink()
            
            # Delete uploaded file
            upload_path = Path(metadata.get("upload_path", ""))
            if upload_path.exists():
                upload_path.unlink()
            
            # Delete metadata
            metadata_path = self.metadata_dir / f"{audiobook_id}.json"
            if metadata_path.exists():
                metadata_path.unlink()
            
            logger.info(f"✅ Audiobook deleted: {audiobook_id}")
            return True
        
        except Exception as e:
            logger.error(f"Error deleting audiobook: {e}")
            return False


# Singleton instance
_audiobook_service: Optional[AudiobookService] = None


def get_audiobook_service() -> AudiobookService:
    """Get or create audiobook service instance"""
    global _audiobook_service
    if _audiobook_service is None:
        _audiobook_service = AudiobookService()
    return _audiobook_service
